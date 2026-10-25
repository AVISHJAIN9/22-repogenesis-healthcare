#!/usr/bin/env python3
import csv
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK_PATHS = [
    Path('/Users/avishjain/avishjain9-repos/24-repogenesis-healthcare/24_RepoGenesis_Healthcare_ULTIMATE_150_Master.xlsx'),
    Path('/Users/avishjain/avishjain9-repos/24-repogenesis-healthcare/audit/24_RepoGenesis_Healthcare_ULTIMATE_150_Master.xlsx'),
]


def row_flag(value):
    if value is None:
        return 0
    if isinstance(value, str):
        value = value.strip()
    return 1 if value in ('✓', 'TRUE', 'True', 'true', 'Y', 'yes', 1, True) else 0


def main():
    workbook_path = next((p for p in WORKBOOK_PATHS if p.exists()), None)
    if workbook_path is None:
        raise FileNotFoundError('Workbook not found.')

    wb = load_workbook(workbook_path, data_only=True)
    sheet = wb['Detailed Analysis']

    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        feature = (row[0] or '').strip()
        if not feature:
            continue
        records.append({
            'feature': feature,
            'FE': row_flag(row[1]),
            'BE': row_flag(row[2]),
            'DB': row_flag(row[3]),
            'AI': row_flag(row[4]),
            'priority': (row[5] or '').strip() if len(row) > 5 else '',
            'area': (row[6] or '').strip() if len(row) > 6 else '',
            'dependencies': (row[7] or '').strip() if len(row) > 7 else '',
        })

    if len(records) != 150:
        raise ValueError(f'Expected 150 feature rows, found {len(records)}')

    output = Path('/Users/avishjain/avishjain9-repos/24-repogenesis-healthcare/audit/spec.csv')
    with output.open('w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['id', 'feature', 'area', 'priority', 'FE', 'BE', 'DB', 'AI', 'dependencies'])
        for idx, r in enumerate(records, start=1):
            writer.writerow([
                f'F{idx:03d}',
                r['feature'],
                r['area'],
                r['priority'],
                r['FE'],
                r['BE'],
                r['DB'],
                r['AI'],
                r['dependencies'],
            ])

    print(f'WROTE {output}')
    print(f'total_rows={len(records)} FE={sum(r["FE"] for r in records)} BE={sum(r["BE"] for r in records)} DB={sum(r["DB"] for r in records)} AI={sum(r["AI"] for r in records)}')

    from collections import Counter
    print('area_counts=' + str(dict(sorted(Counter(r['area'] for r in records).items()))))
    print('priority_counts=' + str(dict(sorted(Counter(r['priority'] for r in records).items()))))


if __name__ == '__main__':
    main()
